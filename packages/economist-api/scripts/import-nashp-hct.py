#!/usr/bin/env python3
"""Offline importer for NASHP Hospital Cost Tool workbooks.

This script is intentionally an offline data-prep utility. The production
Economist API reads the generated JSON artifact from TypeScript/Node.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl
from docx import Document


SOURCE_URL = "https://nashp.org/hospital-cost-tool-and-resources/"

METRICS: list[dict[str, str]] = [
    {
        "slug": "net_patient_revenue",
        "column": "Net Patient Revenue",
        "label": "Net patient revenue",
        "format": "currency",
        "aggregation": "sum",
    },
    {
        "slug": "operating_expenses",
        "column": "Operating Expenses",
        "label": "Operating expenses",
        "format": "currency",
        "aggregation": "sum",
    },
    {
        "slug": "hospital_operating_costs",
        "column": "Hospital Operating Costs",
        "label": "Hospital operating costs",
        "format": "currency",
        "aggregation": "sum",
    },
    {
        "slug": "hospital_expenses_inclusive",
        "column": "Hospital Expenses (Inclusive of All Services)",
        "label": "Hospital expenses inclusive of all services",
        "format": "currency",
        "aggregation": "sum",
    },
    {
        "slug": "direct_patient_care_labor_cost",
        "column": "Direct Patient Care Labor Cost",
        "label": "Direct patient care labor cost",
        "format": "currency",
        "aggregation": "sum",
    },
    {
        "slug": "contracted_labor_cost",
        "column": "Direct Patient Care Contracted Labor Cost",
        "label": "Direct patient care contracted labor cost",
        "format": "currency",
        "aggregation": "sum",
    },
    {
        "slug": "total_drug_costs",
        "column": "Total Drug Costs",
        "label": "Total drug costs",
        "format": "currency",
        "aggregation": "sum",
    },
    {
        "slug": "capital_related_costs",
        "column": "Capital Related Costs",
        "label": "Capital related costs",
        "format": "currency",
        "aggregation": "sum",
    },
    {
        "slug": "other_hospital_operating_costs",
        "column": "Other Hospital Operating Costs",
        "label": "Other hospital operating costs",
        "format": "currency",
        "aggregation": "sum",
    },
    {
        "slug": "adjusted_patient_discharges",
        "column": "Adjusted Patient Discharges",
        "label": "Adjusted patient discharges",
        "format": "count",
        "aggregation": "sum",
    },
    {
        "slug": "bed_size",
        "column": "Bed Size",
        "label": "Bed size",
        "format": "count",
        "aggregation": "sum",
    },
    {
        "slug": "inpatient_occupancy",
        "column": "Inpatient Occupancy",
        "label": "Inpatient occupancy",
        "format": "percent",
        "aggregation": "average",
    },
    {
        "slug": "net_profit_margin",
        "column": "Net Profit Margin",
        "label": "Net profit margin",
        "format": "percent",
        "aggregation": "average",
    },
    {
        "slug": "operating_profit_margin",
        "column": "Operating Profit Margin",
        "label": "Operating profit margin",
        "format": "percent",
        "aggregation": "average",
    },
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build compact NASHP HCT trend data for the Economist API.")
    parser.add_argument("--hct-xlsx", required=True, type=Path, help="NASHP HCT workbook, preferably the all-years file.")
    parser.add_argument("--definitions-docx", required=True, type=Path, help="NASHP variable definitions DOCX.")
    parser.add_argument("--zip-cbsa-csv", type=Path, help="Optional ZIP-to-CBSA crosswalk CSV.")
    parser.add_argument("--out", required=True, type=Path, help="Output JSON artifact path.")
    parser.add_argument("--source-url", default=SOURCE_URL, help="Official NASHP source page URL checked before import.")
    parser.add_argument("--source-page-updated-on", help="Updated On date read from the official NASHP source page.")
    parser.add_argument("--source-checked-at", help="UTC timestamp for the source freshness check.")
    parser.add_argument("--source-page-sha256", help="SHA-256 hash of the fetched official NASHP source page HTML.")
    parser.add_argument("--run-id", help="Operator or scheduler run identifier for audit/debugging.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    zip_cbsa = load_zip_cbsa(args.zip_cbsa_csv) if args.zip_cbsa_csv else {}
    workbook = openpyxl.load_workbook(args.hct_xlsx, read_only=True, data_only=True)
    sheet = resolve_data_sheet(workbook)
    rows = sheet.iter_rows(min_row=1, values_only=True)
    headers = [normalize_header(value) for value in next(rows)]
    indexes = require_columns(headers)

    hospitals: list[list[Any]] = []
    hospital_index: dict[str, int] = {}
    hospital_latest_year: dict[str, int] = {}
    health_systems: list[list[str]] = []
    health_system_index: dict[str, int] = {}
    msas: list[list[str]] = []
    msa_index: dict[str, int] = {}
    observations: list[list[Any]] = []
    years: set[int] = set()
    msa_observations = 0

    for raw in rows:
        year = to_int(raw[indexes["Year"]])
        ccn = normalize_ccn(raw[indexes["CCN#"]])
        if year is None or not ccn:
            continue

        zip_code = normalize_zip(raw[indexes["Zip Code"]])
        cbsa = zip_cbsa.get(zip_code)
        health_system_id = normalize_optional_text(raw[indexes["Health System ID"]])
        health_system_name = normalize_optional_text(raw[indexes["Health System"]])
        health_idx = intern_health_system(health_system_id, health_system_name, health_system_index, health_systems)
        msa_idx = intern_msa(cbsa, msa_index, msas) if cbsa else None
        if msa_idx is not None and msas[msa_idx][0]:
            msa_observations += 1

        hosp_idx = intern_hospital(
            ccn=ccn,
            year=year,
            row=raw,
            indexes=indexes,
            zip_code=zip_code,
            health_idx=health_idx,
            msa_idx=msa_idx,
            hospital_index=hospital_index,
            hospital_latest_year=hospital_latest_year,
            hospitals=hospitals,
        )
        metric_values = [to_number(raw[indexes[metric["column"]]]) for metric in METRICS]
        observations.append([hosp_idx, year, *metric_values])
        years.add(year)

    workbook.close()

    payload = {
        "schemaVersion": 1,
        "source": {
            "name": "NASHP Hospital Cost Tool",
            "url": args.source_url,
            "releaseLabel": infer_release_label(args.hct_xlsx.name),
            "pageUpdatedOn": args.source_page_updated_on,
            "checkedAt": args.source_checked_at,
            "pageSha256": args.source_page_sha256,
            "runId": args.run_id,
            "workbook": source_file(args.hct_xlsx),
            "definitions": {
                **source_file(args.definitions_docx),
                "lastUpdated": read_definitions_last_updated(args.definitions_docx),
            },
            "zipCbsaCrosswalk": source_file(args.zip_cbsa_csv) if args.zip_cbsa_csv else None,
            "importedAt": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "sheetName": sheet.title,
            "rowCount": len(observations),
            "minYear": min(years),
            "maxYear": max(years),
            "msaCoverage": round(msa_observations / max(1, len(observations)), 6),
        },
        "columns": ["hospitalIndex", "year", *[metric["slug"] for metric in METRICS]],
        "metricDefinitions": METRICS,
        "hospitals": hospitals,
        "healthSystems": health_systems,
        "msas": msas,
        "rows": observations,
    }

    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(payload, separators=(",", ":"), ensure_ascii=False), encoding="utf-8")
    print(json.dumps({
        "out": str(args.out),
        "rows": len(observations),
        "hospitals": len(hospitals),
        "healthSystems": len(health_systems),
        "msas": len(msas),
        "minYear": min(years),
        "maxYear": max(years),
        "msaCoverage": payload["source"]["msaCoverage"],
    }, indent=2))


def resolve_data_sheet(workbook: openpyxl.Workbook) -> Any:
    for sheet in workbook.worksheets:
        if sheet.title.lower().startswith("downloadable"):
            return sheet
    raise ValueError("No NASHP Downloadable sheet found.")


def normalize_header(value: Any) -> str:
    return str(value or "").strip()


def require_columns(headers: list[str]) -> dict[str, int]:
    required = [
        "CCN#",
        "Year",
        "Hospital Name",
        "Hospital Abbreviated Name",
        "Zip Code",
        "City",
        "State",
        "Health System ID",
        "Health System",
        *[metric["column"] for metric in METRICS],
    ]
    indexes = {header: index for index, header in enumerate(headers)}
    missing = [column for column in required if column not in indexes]
    if missing:
        raise ValueError(f"Missing expected NASHP columns: {', '.join(missing)}")
    return indexes


def load_zip_cbsa(path: Path) -> dict[str, dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return {row["zip"].zfill(5): row for row in csv.DictReader(handle)}


def intern_health_system(
    health_system_id: str | None,
    health_system_name: str | None,
    index: dict[str, int],
    values: list[list[str]],
) -> int | None:
    if not health_system_id and not health_system_name:
        return None
    key = health_system_id or f"name:{health_system_name}"
    if key not in index:
        index[key] = len(values)
        values.append([health_system_id or key, health_system_name or health_system_id or "Unknown health system"])
    return index[key]


def intern_msa(cbsa: dict[str, str], index: dict[str, int], values: list[list[str]]) -> int | None:
    code = normalize_optional_text(cbsa.get("cbsa_code"))
    title = normalize_optional_text(cbsa.get("cbsa_title"))
    cbsa_type = normalize_optional_text(cbsa.get("cbsa_type"))
    if not title:
        return None
    key = code or f"nonmetro:{title}"
    if key not in index:
        index[key] = len(values)
        values.append([code or "", title, cbsa_type or "Unknown"])
    return index[key]


def intern_hospital(
    *,
    ccn: str,
    year: int,
    row: tuple[Any, ...],
    indexes: dict[str, int],
    zip_code: str,
    health_idx: int | None,
    msa_idx: int | None,
    hospital_index: dict[str, int],
    hospital_latest_year: dict[str, int],
    hospitals: list[list[Any]],
) -> int:
    if ccn not in hospital_index:
        hospital_index[ccn] = len(hospitals)
        hospital_latest_year[ccn] = year
        hospitals.append(hospital_record(ccn, row, indexes, zip_code, health_idx, msa_idx))
        return hospital_index[ccn]

    idx = hospital_index[ccn]
    if year >= hospital_latest_year[ccn]:
        hospitals[idx] = hospital_record(ccn, row, indexes, zip_code, health_idx, msa_idx)
        hospital_latest_year[ccn] = year
    return idx


def hospital_record(
    ccn: str,
    row: tuple[Any, ...],
    indexes: dict[str, int],
    zip_code: str,
    health_idx: int | None,
    msa_idx: int | None,
) -> list[Any]:
    return [
        ccn,
        normalize_text(row[indexes["Hospital Name"]]),
        normalize_optional_text(row[indexes["Hospital Abbreviated Name"]]),
        normalize_text(row[indexes["City"]]),
        normalize_text(row[indexes["State"]]),
        zip_code,
        health_idx,
        msa_idx,
    ]


def normalize_ccn(value: Any) -> str | None:
    text = normalize_optional_text(value)
    if not text:
        return None
    return text.split(".")[0].zfill(6)


def normalize_zip(value: Any) -> str:
    text = normalize_optional_text(value)
    if not text:
        return ""
    return text.split(".")[0].zfill(5)[:5]


def normalize_optional_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.upper() in {"NA", "N/A", "NONE", "NULL", "."}:
        return None
    return text


def normalize_text(value: Any) -> str:
    return normalize_optional_text(value) or "Unknown"


def to_int(value: Any) -> int | None:
    number = to_number(value)
    return int(number) if number is not None else None


def to_number(value: Any) -> float | int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        if not math.isfinite(float(value)):
            return None
        return int(value) if float(value).is_integer() else round(float(value), 6)
    text = normalize_optional_text(value)
    if not text:
        return None
    try:
        number = float(text.replace(",", ""))
    except ValueError:
        return None
    if not math.isfinite(number):
        return None
    return int(number) if number.is_integer() else round(number, 6)


def source_file(path: Path | None) -> dict[str, Any] | None:
    if not path:
        return None
    return {
        "fileName": path.name,
        "sha256": sha256(path),
        "sizeBytes": path.stat().st_size,
    }


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def infer_release_label(file_name: str) -> str:
    coverage = re.search(r"NASHP\s+(.+?)\s+HCT\s+Data", file_name, flags=re.IGNORECASE)
    release = re.search(r"HCT\s+Data\s+(20\d{2})\s+([A-Za-z]{3,9})", file_name, flags=re.IGNORECASE)
    if coverage and release:
        return f"{coverage.group(1)} HCT data, {release.group(1)} {release.group(2)}"
    if release:
        return f"HCT data, {release.group(1)} {release.group(2)}"
    return file_name


def read_definitions_last_updated(path: Path) -> str | None:
    doc = Document(path)
    for paragraph in doc.paragraphs[:20]:
        text = paragraph.text.strip()
        if text.lower().startswith("last updated:"):
            return text.split(":", 1)[1].strip()
    return None


if __name__ == "__main__":
    main()
